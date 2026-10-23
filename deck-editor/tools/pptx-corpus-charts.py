"""Adds native charts to the corpus deck's "Charts" slide (called by
pptx-corpus.ps1, which cannot make charts through COM without a window).

    py -3 tools/pptx-corpus-charts.py <deck.pptx>      needs python-pptx (dev only)
"""
import sys

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.util import Pt

path = sys.argv[1]
prs = Presentation(path)
slide = next(s for s in prs.slides if s.shapes.title is not None and s.shapes.title.text == "Charts")

cd = CategoryChartData()
cd.categories = ["Q1", "Q2", "Q3", "Q4"]
cd.add_series("2025", (4.2, 4.8, 5.1, 5.9))
cd.add_series("2026", (5.0, 5.6, 6.3, 7.1))
gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Pt(40), Pt(120), Pt(440), Pt(380), cd)
ch = gf.chart
ch.has_title = True
ch.chart_title.text_frame.text = "Revenue by quarter (M)"
ch.has_legend = True
ch.legend.position = XL_LEGEND_POSITION.BOTTOM
ch.legend.include_in_layout = False

cd2 = CategoryChartData()
cd2.categories = ["Enterprise", "Self-serve", "Education", "Other"]
cd2.add_series("Share", (46, 31, 15, 8))
gf2 = slide.shapes.add_chart(XL_CHART_TYPE.PIE, Pt(500), Pt(120), Pt(420), Pt(380), cd2)
ch2 = gf2.chart
ch2.has_title = True
ch2.chart_title.text_frame.text = "Share by segment"
ch2.plots[0].has_data_labels = True
ch2.plots[0].data_labels.number_format = '0"%"'
ch2.plots[0].data_labels.number_format_is_linked = False

# A line chart on the same slide would crowd it; the two above cover bars and pies.
prs.save(path)
print("charts added to", path)
